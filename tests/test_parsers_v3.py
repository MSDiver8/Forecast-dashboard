"""Parsers of the sources behind US real GDP, UK CPI and UK real GDP (added 13 October 2026)."""

import json
from datetime import date

import pytest
from openpyxl import load_workbook

from fdash.ingest import bea_nipa, boe_mpr_uk, cbo_econ, fannie_mae, obr_efo, ons, spf_philly, td_economics
from fdash.ingest.base import Release, SourceStructureError, projection_kind


def release(source_id: str, vintage: date, **extra) -> Release:
    return Release(f"{source_id}_test", source_id, "test", vintage, "http://example", "x", extra)


def test_projection_kind_before_release_quarter_is_estimate():
    assert projection_kind("2026-Q2", date(2026, 7, 30), True) == "estimate"
    assert projection_kind("2026-Q3", date(2026, 7, 30), True) == "forecast"
    assert projection_kind("2026-Q3", date(2026, 7, 30), False) == "actual"


def test_bea_table_111(fixtures):
    published, points = bea_nipa.read_file(fixtures / "bea_nipa" / "Section1All_T10101-Q.xlsx")
    assert published == date(2026, 9, 30)
    assert points[0] == ("1947-Q2", -1.0)
    assert dict(points)["2026-Q2"] == 2.2 and points[-1][0] == "2026-Q2"


def test_ons_quarters_and_release_date(fixtures):
    cpi = json.loads((fixtures / "ons" / "d7g7_mm23.json").read_text(encoding="utf-8"))
    gdp = json.loads((fixtures / "ons" / "ihyr_qna.json").read_text(encoding="utf-8"))
    assert ons.release_date(cpi) == date(2026, 9, 16)  # 2026-09-15T23:00Z is midnight in London
    assert ons.read_quarters(cpi, "D7G7")[-1] == ("2026-Q2", 2.8)
    assert ons.read_quarters(gdp, "IHYR")[-1] == ("2026-Q2", 1.4)
    with pytest.raises(SourceStructureError):
        ons.read_quarters(gdp, "D7G7")


def test_boe_history_rows_and_shading(fixtures):
    history = boe_mpr_uk.read_history(
        load_workbook(fixtures / "boe_mpr_uk" / "databank-july-2026-history.xlsx", data_only=True)
    )
    assert len(history) == 90 and min(history) == date(2004, 2, 1)
    assert date(2026, 4, 1) not in history  # no central projection in the April 2026 Report
    latest = history[date(2026, 8, 1)]
    cpi = [(p, v) for p, v, projected in latest["cpi_yoy_q"] if projected]
    assert cpi[0] == ("2026-Q3", 2.9) and cpi[-1][0] == "2029-Q3" and len(cpi) == 13
    assert [(p, v) for p, v, projected in latest["cpi_yoy_q"] if not projected][-1] == ("2026-Q2", 2.8)
    gdp = [p for p, _, projected in latest["gdp_yoy_q"] if projected]
    assert gdp[0] == "2026-Q2" and len(gdp) == 14  # the backcast of the previous quarter is shaded too
    assert "gdp_yoy_q" not in history[date(2004, 2, 1)]


def test_boe_row_dates():
    exact = [date(2025, 11, 6), date(2026, 2, 5), date(2026, 4, 30), date(2026, 7, 30)]
    assert boe_mpr_uk.match_exact(date(2026, 8, 1), exact) == date(2026, 7, 30)
    assert boe_mpr_uk.match_exact(date(2025, 11, 1), exact) == date(2025, 11, 6)
    assert boe_mpr_uk.match_exact(date(2025, 8, 1), exact) is None


def test_obr_tables_and_kinds(fixtures):
    raw = fixtures / "obr_efo" / "economy-march-2026.xlsx"
    cpi, growth, outturn = obr_efo.read_tables(load_workbook(raw, data_only=True))
    assert outturn == "2025-12"
    assert dict(cpi)["2026-Q2"] == pytest.approx(2.137, abs=0.001)
    assert dict(growth)["2026-Q3"] == pytest.approx(1.215, abs=0.001)
    parsed = obr_efo.ObrEfo().parse(raw, release("obr_efo", date(2026, 3, 3)))
    kinds = {(o.indicator_id, o.target_period): o.kind for o in parsed.observations}
    assert kinds[("cpi_yoy_q", "2025-Q4")] == "actual" and kinds[("cpi_yoy_q", "2026-Q1")] == "forecast"
    assert kinds[("gdp_yoy_q", "2025-Q4")] == "estimate" and kinds[("gdp_yoy_q", "2026-Q1")] == "forecast"
    html = (fixtures / "obr_efo" / "efo-march-2026-excerpt.html").read_text(encoding="utf-8")
    assert obr_efo.published_date(html) == date(2026, 3, 3)


def test_obr_tells_an_html_answer_from_the_tables(fixtures, tmp_path):
    page = tmp_path / "economy.xlsx"
    page.write_text(
        "<!DOCTYPE html><title>No Access - Office for Budget Responsibility</title>", encoding="utf-8"
    )
    assert not obr_efo.is_xlsx(page)
    assert obr_efo.is_xlsx(fixtures / "obr_efo" / "economy-march-2026.xlsx")


@pytest.mark.parametrize(
    "month,quarter",
    [("2025-12", "2025-Q4"), ("2025-08", "2025-Q2"), ("2025-01", "2024-Q4"), ("2024-09", "2024-Q3")],
)
def test_obr_last_complete_quarter(month, quarter):
    assert obr_efo.last_complete_quarter(month) == quarter


def test_cbo_files_and_kinds(fixtures):
    catalog = json.loads((fixtures / "cbo_econ" / "catalog-excerpt.json").read_text(encoding="utf-8"))
    assert [v for v, _ in cbo_econ.quarterly_files(catalog)] == [
        "2024-02",
        "2024-06",
        "2025-01",
        "2025-09",
        "2026-02",
    ]
    raw = fixtures / "cbo_econ" / "quarterly_2026-02.csv"
    parsed = cbo_econ.CboEcon().parse(raw, release("cbo_econ", date(2026, 2, 1)))
    kinds = {o.target_period: o.kind for o in parsed.observations}
    assert kinds["2025-Q2"] == "actual" and kinds["2025-Q4"] == "estimate" and kinds["2026-Q1"] == "forecast"
    assert parsed.data_cutoff == "2025-Q2"
    assert {o.indicator_id for o in parsed.observations} == {"gdp_qoq_saar"}


def test_spf_release_dates_and_survey(fixtures):
    dates = spf_philly.release_dates(
        (fixtures / "spf_philly" / "spf-release-dates.txt").read_text(encoding="utf-8")
    )
    assert dates[(2026, 3)] == date(2026, 8, 14)
    assert dates[(2026, 1)] == date(2026, 3, 6)  # delayed survey marked with asterisks
    assert dates[(1990, 2)] == date(1990, 8, 31)
    parsed = spf_philly.SpfPhilly().parse(
        fixtures / "spf_philly" / "mediangrowth.xlsx",
        release("spf_philly", date(2026, 8, 14), survey="2026-Q3"),
    )
    values = [(o.target_period, round(o.value, 2)) for o in parsed.observations]
    assert values == [
        ("2026-Q3", 2.46),
        ("2026-Q4", 2.27),
        ("2027-Q1", 2.12),
        ("2027-Q2", 2.12),
        ("2027-Q3", 2.24),
    ]


def test_fannie_mae_table(fixtures):
    published, points = fannie_mae.read_pdf(fixtures / "fannie_mae" / "economic-forecast-092026.pdf")
    assert published == date(2026, 9, 11)
    assert [p for p, _, _ in points] == [f"{y}-Q{q}" for y in (2025, 2026, 2027) for q in (1, 2, 3, 4)]
    assert [p for p, _, projected in points if projected][0] == "2026-Q3"
    assert dict((p, v) for p, v, _ in points)["2026-Q3"] == 2.4
    raw = fixtures / "fannie_mae" / "economic-forecast-102024.pdf"
    parsed = fannie_mae.FannieMae().parse(raw, release("fannie_mae", date(2024, 10, 10)))
    kinds = {o.target_period: o.kind for o in parsed.observations}
    assert kinds["2024-Q2"] == "actual" and kinds["2024-Q3"] == "estimate" and kinds["2024-Q4"] == "forecast"


class _Answer:
    def __init__(self, status: int, content_type: str = "text/html", text: str = "") -> None:
        self.status_code, self.headers, self.text, self.content = (
            status,
            {"content-type": content_type},
            text,
            b"",
        )

    def raise_for_status(self) -> None:
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


def test_fannie_mae_date_typo_takes_the_year_of_the_title(fixtures):
    # "Economic Forecast: January 2021", dated "January 11, 2020" in the text
    published, points = fannie_mae.read_pdf(fixtures / "fannie_mae" / "economic-forecast-011521.pdf")
    assert published == date(2021, 1, 11)
    assert points[0][0] == "2020-Q1" and points[-1][0] == "2022-Q4"


def test_fannie_mae_archive_links(fixtures):
    html = (fixtures / "fannie_mae" / "forecast-monthly-archive-excerpt.html").read_text(encoding="utf-8")
    links = fannie_mae.archive_links(html)
    assert sorted(links) == [(2020, 9), (2022, 1), (2022, 7), (2026, 9), (2026, 10)]
    assert links[(2020, 9)].endswith("/economic-forecast-091520")  # month, day and year
    assert links[(2022, 1)].endswith("/economic-forecast-011922pdf")
    assert links[(2026, 10)].endswith("/economic-forecast-102026")  # month and year


def test_fannie_mae_reports_a_blocked_check(fixtures, monkeypatch):
    # a 403 instead of the PDF must not pass for "nothing new"
    html = (fixtures / "fannie_mae" / "forecast-monthly-archive-excerpt.html").read_text(encoding="utf-8")
    monkeypatch.setattr(fannie_mae, "saved_releases", lambda source_id: {})
    monkeypatch.setattr(
        fannie_mae,
        "http_get",
        lambda url, **kwargs: _Answer(200, text=html) if url == fannie_mae.ARCHIVE else _Answer(403),
    )
    with pytest.raises(SourceStructureError, match="answer 403"):
        fannie_mae.FannieMae().list_releases()


def test_td_table(fixtures):
    published, points = td_economics.read_pdf(
        fixtures / "td_economics" / "QEF_September2026_US-pages-1-5.pdf"
    )
    assert published == date(2026, 9, 17)
    projected = {p: flag for p, _, flag in points}
    assert projected["2026-Q2"] is False and projected["2026-Q3"] is True
    assert projected["2027-Q1"] is True  # printed without F in this report
    assert dict((p, v) for p, v, _ in points)["2026-Q3"] == 2.8
    with pytest.raises(SourceStructureError):
        td_economics.read_pdf(fixtures / "fannie_mae" / "economic-forecast-092026.pdf")


def test_td_links(fixtures):
    html = (fixtures / "td_economics" / "us-quarterly-economic-forecast-excerpt.html").read_text(
        encoding="utf-8"
    )
    year, month, url = td_economics.newest_link(html)
    assert (year, month) == (2026, 9) and url.endswith("/2026-sep/QEF_September2026_US.pdf")
    assert any(u.endswith("/2026-june/QEF_June2026_US.pdf") for u in td_economics.candidates(2026, 6))
    assert any(u.endswith("/2024-dec/qefdec2024_us.pdf") for u in td_economics.candidates(2024, 12))
