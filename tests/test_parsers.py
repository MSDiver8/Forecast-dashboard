"""Parsers against saved source responses (tests/fixtures/)."""

from datetime import date

import pytest
from openpyxl import load_workbook

from fdash.ingest import boe_mpr, cbr_stats, cbr_survey, ecb_mpd, eia_brent_spot, eia_steo, imf_weo, wb_cmo
from fdash.ingest.base import Release


def release(source_id: str, vintage: date, **extra) -> Release:
    return Release(f"{source_id}_test", source_id, "test", vintage, "http://example", "x", extra)


def test_eia_release_page_lists_two_years(fixtures):
    html = (fixtures / "eia_steo" / "outlook.html").read_text(encoding="utf-8")
    found = eia_steo.parse_release_page(html, date(2026, 10, 7))
    assert found[-1] == ("oct26", date(2026, 10, 6))
    assert ("jun25", date(2025, 6, 10)) in found  # dates without a leading zero
    assert len(found) == 26


def test_eia_steo_brent_months(fixtures):
    parsed = eia_steo.EiaSteo().parse(
        fixtures / "eia_steo" / "oct26_base.xlsx", release("eia_steo", date(2026, 10, 6))
    )
    by_period = {o.target_period: o for o in parsed.observations}
    assert parsed.data_cutoff == "2026-09"
    assert by_period["2022-01"].value == pytest.approx(86.51)
    assert by_period["2026-09"].kind == "actual" and by_period["2026-10"].kind == "forecast"


def test_eia_brent_spot_monthly_and_annual(fixtures):
    html = (fixtures / "eia_brent_spot" / "rbrte.html").read_text(encoding="utf-8")
    assert eia_brent_spot.release_date(html) == date(2026, 9, 30)
    monthly = dict(eia_brent_spot.read_series(fixtures / "eia_brent_spot" / "RBRTEm.xls", "M"))
    annual = dict(eia_brent_spot.read_series(fixtures / "eia_brent_spot" / "RBRTEa.xls", "A"))
    assert monthly["1987-05"] == pytest.approx(18.58)
    assert annual["2021"] == pytest.approx(70.86)


@pytest.mark.parametrize("name,brent_2026", [("CMO-April-2026-Forecasts.xlsx", 86.0)])
def test_wb_cmo_new_layout(fixtures, name, brent_2026):
    rows = list(
        load_workbook(fixtures / "wb_cmo" / name, read_only=True)["Forecast"].iter_rows(
            max_col=40, values_only=True
        )
    )
    points, cutoff = wb_cmo.parse_forecast_sheet(rows)
    values = {y: (v, f) for y, v, f in points}
    assert values["2026"] == (brent_2026, True)
    assert cutoff == "2025"
    assert "2026" in values and len(values) == 4  # the trailing "% change" column is not a year


def test_wb_cmo_old_layout(fixtures):
    rows = list(
        load_workbook(fixtures / "wb_cmo" / "CMO-April-2022-forecasts.xlsx", read_only=True)[
            "Forecast"
        ].iter_rows(max_col=40, values_only=True)
    )
    points, _ = wb_cmo.parse_forecast_sheet(rows)
    assert dict((y, v) for y, v, _ in points)["2022"] == 100.0


def test_wb_archive_starts_with_brent_row(fixtures):
    html = (fixtures / "wb_cmo" / "price-forecasts.html").read_text(encoding="utf-8")
    found = wb_cmo.parse_archive(html)
    assert found[0][:2] == (2022, 4) and found[-1][:2] == (2026, 4)


def test_boe_assumption_is_q4_average(fixtures):
    sheet = load_workbook(
        fixtures / "boe_mpr" / "databank-july-2026-other-assumptions.xlsx", read_only=True
    ).active
    points, cutoff = boe_mpr.parse_assumptions(list(sheet.iter_rows(values_only=True)))
    assert points[0] == ("2024-Q4", pytest.approx(74.667))
    assert [p for p, _ in points] == ["2024-Q4", "2025-Q4", "2026-Q4", "2027-Q4", "2028-Q4"]
    assert cutoff == "20 July 2026"


def test_boe_unpublished_report_is_skipped(fixtures):
    html = (fixtures / "boe_mpr" / "november-2026-unpublished.html").read_text(encoding="utf-8")
    assert boe_mpr.parse_release_page(html) is None
    published = boe_mpr.parse_release_page(
        (fixtures / "boe_mpr" / "july-2026.html").read_text(encoding="utf-8")
    )
    assert published[0] == date(2026, 7, 30) and published[1].endswith(".zip")


def test_ecb_publication_dates(fixtures):
    html = (fixtures / "ecb_mpd" / "projections-index.html").read_text(encoding="utf-8")
    dates = ecb_mpd.publication_dates(html)
    assert dates[(2026, 6)] == date(2026, 6, 11)
    assert dates[(2026, 3)] == date(2026, 3, 19)


def test_ecb_mpd_oil_and_macro(fixtures):
    oil = ecb_mpd.EcbMpd().parse(
        fixtures / "ecb_mpd" / "pou-S26.csv", release("ecb_mpd", date(2026, 9, 10), exercise="S26")
    )
    annual = {o.target_period: o for o in oil.observations if o.frequency == "A"}
    assert annual["2026"].value == pytest.approx(89.5) and annual["2026"].kind == "forecast"
    assert annual["2025"].kind == "estimate"
    macro = ecb_mpd.EcbMpdMacro().parse(
        fixtures / "ecb_mpd" / "macro-U2-S26.csv", release("ecb_mpd_macro", date(2026, 9, 10), exercise="S26")
    )
    got = {(o.indicator_id, o.target_period): o.value for o in macro.observations}
    assert got[("gdp_growth", "2026")] == pytest.approx(0.9)
    assert got[("cpi_avg", "2027")] == pytest.approx(2.5)
    assert got[("unemployment_avg", "2028")] == pytest.approx(5.9)
    assert {o.area_id for o in macro.observations} == {"EA"}


def test_cbr_survey_latest_round(fixtures):
    raw = fixtures / "cbr_survey" / "full.xlsx"
    parsed = cbr_survey.CbrSurvey().parse(
        raw, release("cbr_survey", date(2026, 9, 1), survey_month="2026-09-01")
    )
    cpi = {o.target_period: o for o in parsed.observations if o.indicator_id == "cpi_dec"}
    assert cpi["2026"].value == pytest.approx(6.6)
    assert cpi["2026"].lower <= cpi["2026"].value <= cpi["2026"].upper
    rates = {o.target_period: o.value for o in parsed.observations if o.indicator_id == "key_rate_avg"}
    assert rates["2026"] == pytest.approx(14.5, abs=0.05)


def test_imf_weo_kinds(fixtures):
    parsed = imf_weo.ImfWeo().parse(
        fixtures / "imf_weo" / "WEO_2025_OCT_VINTAGE_sample.csv", release("imf_weo", date(2025, 10, 14))
    )
    ru = {
        o.target_period: o
        for o in parsed.observations
        if o.indicator_id == "gdp_growth" and o.area_id == "RU"
    }
    assert ru["2024"].kind == "estimate" and ru["2025"].kind == "forecast"
    assert ru["2025"].value == pytest.approx(0.609)


def test_cbr_stats_matches_published_figures(fixtures):
    inflation = cbr_stats.read_inflation(fixtures / "cbr_stats" / "infl.xlsx")
    assert inflation["2025-12"] == pytest.approx(5.59, abs=0.01)
    rates = {date(2024, 1, 1): 16.0, date(2024, 7, 29): 18.0, date(2024, 10, 28): 21.0}
    avg = cbr_stats.annual_average_calendar_days(rates, date(2025, 1, 10))
    expected = (16 * 210 + 18 * 91 + 21 * 65) / 366
    assert avg["2024"] == pytest.approx(expected)
    soap = cbr_stats.read_key_rate(fixtures / "cbr_stats" / "keyrate-soap-response.xml")
    assert soap[date(2026, 1, 9)] == 16.0
