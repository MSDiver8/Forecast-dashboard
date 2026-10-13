"""OBR Economic and fiscal outlook, detailed forecast tables (economy): UK CPI and real GDP, quarterly."""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from fdash.ingest.base import (
    Connector,
    Observation,
    Parsed,
    Release,
    SourceStructureError,
    archive_start,
    client,
    quarter_of,
)
from fdash.ingest.registry import register

BASE = "https://obr.uk"
MONTHS = [
    "january",
    "february",
    "march",
    "april",
    "may",
    "june",
    "july",
    "august",
    "september",
    "october",
    "november",
    "december",
]
PERIOD = re.compile(r"^(\d{4})Q([1-4])$")
OUTTURN = re.compile(r"CPI\b.*outturn data up to and including (\w+ \d{4})")
PUBLISHED = re.compile(r'"datePublished":"(\d{4}-\d{2}-\d{2})')


def page_url(year: int, month: int) -> str:
    return f"{BASE}/efo/economic-and-fiscal-outlook-{MONTHS[month - 1]}-{year}/"


def tables_url(year: int, month: int) -> str:
    slug = f"{MONTHS[month - 1]}-{year}-economic-and-fiscal-outlook"
    return f"{BASE}/download/{slug}-detailed-forecast-tables-economy/"


def published_date(html: str) -> date:
    match = PUBLISHED.search(html)
    if not match:
        raise SourceStructureError("EFO page: datePublished not found")
    return date.fromisoformat(match.group(1))


def _column(sheet, title: str, rows: range) -> int:
    for r in rows:
        for c in range(1, sheet.max_column + 1):
            if str(sheet.cell(r, c).value or "").strip() == title:
                return c
    raise SourceStructureError(f"sheet {sheet.title}: column {title!r} not found")


def _quarters(sheet, column: int) -> list[tuple[str, float]]:
    points = []
    for r in range(5, sheet.max_row + 1):
        match = PERIOD.match(str(sheet.cell(r, 2).value or ""))
        value = sheet.cell(r, column).value
        if match and isinstance(value, int | float) and not isinstance(value, bool):
            points.append((f"{match.group(1)}-Q{match.group(2)}", float(value)))
    return points


def read_tables(workbook) -> tuple[list[tuple[str, float]], list[tuple[str, float]], str]:
    """CPI (% on a year earlier), real GDP growth on a year earlier from levels, last CPI outturn month."""
    for name in ("1.1", "1.7"):
        if name not in workbook.sheetnames:
            raise SourceStructureError(f"sheet {name} not found")
    inflation, gdp = workbook["1.7"], workbook["1.1"]
    if "Inflation" not in str(inflation.cell(2, 2).value) or "chain-linked volumes" not in str(
        gdp.cell(2, 2).value
    ):
        raise SourceStructureError("tables 1.7 Inflation / 1.1 chain-linked volumes not found")
    cpi = _quarters(inflation, _column(inflation, "CPI", range(3, 6)))
    note = " ".join(
        str(c.value) for row in inflation.iter_rows(min_row=100) for c in row if isinstance(c.value, str)
    )
    outturn = OUTTURN.search(note)
    if not outturn:
        raise SourceStructureError("1.7: CPI outturn note not found")
    last = datetime.strptime(outturn.group(1), "%B %Y")
    levels = dict(_quarters(gdp, _column(gdp, "Real GDP", range(2, 6))))
    growth = []
    for period, level in levels.items():
        previous = f"{int(period[:4]) - 1}{period[4:]}"
        if previous in levels:
            growth.append((period, 100 * (level / levels[previous] - 1)))
    return cpi, growth, f"{last.year:04d}-{last.month:02d}"


def last_complete_quarter(month: str) -> str:
    year, m = int(month[:4]), int(month[5:7])
    if m % 3:
        m -= m % 3
        if m == 0:
            year, m = year - 1, 12
    return f"{year:04d}-Q{m // 3}"


class ObrEfo(Connector):
    source_id = "obr_efo"
    name = "Economic and fiscal outlook: detailed forecast tables, economy"
    organization = "OBR"
    url = f"{BASE}/publications/"
    license = "Open Government Licence v3.0"
    access = "XLSX по ссылке со страницы выпуска EFO"

    def list_releases(self) -> list[Release]:
        start = archive_start(2)
        today = date.today()
        releases = []
        with client() as http:
            year, month = start.year, start.month
            while (year, month) <= (today.year, today.month):
                if http.head(tables_url(year, month)).status_code in (200, 302):
                    page = http.get(page_url(year, month))
                    page.raise_for_status()
                    published = published_date(page.text)
                    releases.append(
                        Release(
                            release_id=f"{self.source_id}_{year}_{month:02d}",
                            source_id=self.source_id,
                            title=f"EFO {month:02d}.{year}",
                            vintage_date=published,
                            url=tables_url(year, month),
                            filename="economy.xlsx",
                        )
                    )
                year, month = (year + 1, 1) if month == 12 else (year, month + 1)
        return releases

    def parse(self, raw: Path, release: Release) -> Parsed:
        cpi, growth, outturn = read_tables(load_workbook(raw, data_only=True))
        release_quarter = quarter_of(release.vintage_date)
        known = last_complete_quarter(outturn)

        def cpi_kind(period: str) -> str:
            if period <= known:
                return "actual"
            return "forecast" if period >= release_quarter else "estimate"

        # The economy tables do not mark which GDP quarters are ONS outturns: before the release
        # quarter the values are kept as estimates, from it on as the forecast.
        observations = [Observation("cpi_yoy_q", "GB", p, "Q", v, cpi_kind(p)) for p, v in cpi] + [
            Observation("gdp_yoy_q", "GB", p, "Q", v, "forecast" if p >= release_quarter else "estimate")
            for p, v in growth
        ]
        return Parsed(observations, data_cutoff=outturn)


register(ObrEfo())
