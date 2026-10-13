"""BEA NIPA Table 1.1.1: real GDP, percent change from the preceding quarter at annual rate (A191RL)."""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from fdash.ingest.base import (
    Connector,
    Observation,
    Parsed,
    Release,
    SourceStructureError,
    copy_into,
    download_temp,
)
from fdash.ingest.registry import register

FILE = "https://apps.bea.gov/national/Release/XLS/Survey/Section1All_xls.xlsx"
SHEET = "T10101-Q"
CODE = "A191RL"
PUBLISHED = re.compile(r"Data published (\w+ \d{1,2}, \d{4})")
PERIOD = re.compile(r"^(\d{4})Q([1-4])$")


def read_table(rows: list[tuple]) -> tuple[date, list[tuple[str, float]]]:
    """Publication date and (YYYY-Qn, value) of line 1 'Gross domestic product' of Table 1.1.1."""
    head = " ".join(str(r[0]) for r in rows[:8] if r and r[0])
    if "Table 1.1.1" not in head or "seasonally adjusted at annual rates" not in head.lower():
        raise SourceStructureError(f"{SHEET}: not Table 1.1.1 at annual rates: {head[:120]!r}")
    published = PUBLISHED.search(head)
    if not published:
        raise SourceStructureError(f"{SHEET}: 'Data published' not found")
    header = next((r for r in rows if r and r[0] == "Line"), None)
    line = next((r for r in rows if r and len(r) > 2 and r[2] == CODE), None)
    if header is None or line is None:
        raise SourceStructureError(f"{SHEET}: header row or series {CODE} not found")
    points = []
    for label, value in zip(header[3:], line[3:], strict=False):
        match = PERIOD.match(str(label or ""))
        if match and isinstance(value, int | float):
            points.append((f"{match.group(1)}-Q{match.group(2)}", float(value)))
    return datetime.strptime(published.group(1), "%B %d, %Y").date(), points


def read_file(path: Path) -> tuple[date, list[tuple[str, float]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SourceStructureError(f"{path.name}: sheet {SHEET} not found")
    return read_table(list(wb[SHEET].iter_rows(values_only=True)))


class BeaNipa(Connector):
    source_id = "bea_nipa"
    name = "NIPA Table 1.1.1: Real GDP, percent change at annual rate (A191RL)"
    organization = "BEA"
    url = "https://www.bea.gov/data/gdp/gross-domestic-product"
    license = "Public domain (U.S. Government)"
    access = "XLSX раздела 1 NIPA по стабильной ссылке, без ключа"
    role = "actual"
    dedupe = True

    def __init__(self) -> None:
        self._latest: Path | None = None

    def _download_latest(self) -> Path:
        if self._latest is None or not self._latest.exists():
            self._latest = download_temp(FILE, "Section1All_xls.xlsx")
        return self._latest

    def list_releases(self) -> list[Release]:
        published, _ = read_file(self._download_latest())
        return [
            Release(
                release_id=f"bea_nipa_{published:%Y_%m_%d}",
                source_id=self.source_id,
                title=f"NIPA {published:%d.%m.%Y}",
                vintage_date=published,
                url=FILE,
                filename="Section1All_xls.xlsx",
            )
        ]

    def fetch(self, release: Release) -> Path:
        return copy_into(self._download_latest(), release)

    def parse(self, raw: Path, release: Release) -> Parsed:
        _, points = read_file(raw)
        observations = [Observation("gdp_qoq_saar", "US", p, "Q", v, "actual") for p, v in points]
        return Parsed(observations, data_cutoff=points[-1][0] if points else None)


register(BeaNipa())
