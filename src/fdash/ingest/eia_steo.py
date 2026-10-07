"""EIA Short-Term Energy Outlook: monthly Brent spot price, one XLSX per release."""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from fdash.ingest.base import Connector, Observation, Parsed, Release, SourceStructureError, http_get
from fdash.ingest.registry import register

PAGE = "https://www.eia.gov/outlooks/steo/outlook.php"
ARCHIVE = "https://www.eia.gov/outlooks/steo/archives/"
MONTHS = {
    m: i
    for i, m in enumerate(
        ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], start=1
    )
}
ROW_DATE = re.compile(r"<td>\s*(\d{1,2}/\d{1,2}/\d{4})\s*</td>")
ROW_FILE = re.compile(r"archives/([a-z]{3}\d{2})_base\.xlsx")
CURRENT_DATE = re.compile(r"Release Date:</strong>\s*([A-Z][a-z]+ \d{1,2}, \d{4})")
SERIES = {"BREPUUS": ("brent_price", "WORLD")}


def parse_release_page(html: str, today: date, years: int = 2) -> list[tuple[str, date]]:
    """(file stem, publication date) for the current release and the archive."""
    found: dict[str, date] = {}
    for row in html.split("<tr")[1:]:
        published, stem = ROW_DATE.search(row), ROW_FILE.search(row)
        if published and stem:
            found[stem.group(1)] = datetime.strptime(published.group(1), "%m/%d/%Y").date()
    current = CURRENT_DATE.search(html)
    if current:
        published = datetime.strptime(current.group(1), "%B %d, %Y").date()
        stem = published.strftime("%b%y").lower()
        found.setdefault(stem, published)
    since = today - timedelta(days=365 * years + 31)
    return sorted(((s, d) for s, d in found.items() if d >= since), key=lambda item: item[1])


class EiaSteo(Connector):
    source_id = "eia_steo"
    name = "Short-Term Energy Outlook"
    organization = "EIA"
    url = PAGE
    license = "Public domain (U.S. Government)"
    access = "XLSX по стабильной ссылке"

    def list_releases(self) -> list[Release]:
        response = http_get(PAGE)
        response.raise_for_status()
        return [
            Release(
                release_id=f"eia_steo_{published:%Y_%m}",
                source_id=self.source_id,
                title=f"STEO {published:%m.%Y}",
                vintage_date=published,
                url=f"{ARCHIVE}{stem}_base.xlsx",
                filename=f"{stem}_base.xlsx",
            )
            for stem, published in parse_release_page(response.text, date.today())
        ]

    def parse(self, raw: Path, release: Release) -> Parsed:
        wb = load_workbook(raw, read_only=True, data_only=True)
        if "Dates" not in wb.sheetnames or "2tab" not in wb.sheetnames:
            raise SourceStructureError(f"{raw.name}: sheets Dates/2tab not found")
        cutoff_code = wb["Dates"].cell(7, 4).value
        if not isinstance(cutoff_code, int):
            raise SourceStructureError(f"{raw.name}: Dates!D7 is not YYYYMM: {cutoff_code!r}")
        cutoff = f"{cutoff_code // 100:04d}-{cutoff_code % 100:02d}"
        sheet = wb["2tab"]
        rows = {sheet.cell(r, 1).value: r for r in range(1, min(sheet.max_row, 200) + 1)}
        observations = []
        for code, (indicator_id, area_id) in SERIES.items():
            if code not in rows:
                raise SourceStructureError(f"{raw.name}: series {code} not found on 2tab")
            year = None
            for col in range(3, sheet.max_column + 1):
                header_year = sheet.cell(3, col).value
                if isinstance(header_year, int | float) and 1990 < header_year < 2100:
                    year = int(header_year)
                month = MONTHS.get(str(sheet.cell(4, col).value).strip())
                value = sheet.cell(rows[code], col).value
                if year is None or month is None or not isinstance(value, int | float):
                    continue
                period = f"{year:04d}-{month:02d}"
                observations.append(
                    Observation(
                        indicator_id,
                        area_id,
                        period,
                        "M",
                        float(value),
                        "actual" if period <= cutoff else "forecast",
                    )
                )
        return Parsed(observations, data_cutoff=cutoff)


register(EiaSteo())
