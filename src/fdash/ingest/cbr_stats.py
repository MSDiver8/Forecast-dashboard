"""Bank of Russia statistics as reference actuals: CPI y/y (December = Dec/Dec) and the key rate."""

from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

from fdash.ingest.base import Connector, Observation, Parsed, Release, SourceStructureError, client, download
from fdash.ingest.registry import register

INFL = (
    "https://www.cbr.ru/Queries/UniDbQuery/DownloadExcel/132934?FromDate=01%2F01%2F2015"
    "&ToDate={to}&posted=False"
)
SOAP = "https://www.cbr.ru/DailyInfoWebServ/DailyInfo.asmx"
SOAP_BODY = (
    '<?xml version="1.0" encoding="utf-8"?>'
    '<soap12:Envelope xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
    'xmlns:xsd="http://www.w3.org/2001/XMLSchema" xmlns:soap12="http://www.w3.org/2003/05/soap-envelope">'
    '<soap12:Body><KeyRateXML xmlns="http://web.cbr.ru/"><fromDate>{start}</fromDate>'
    "<ToDate>{end}</ToDate></KeyRateXML></soap12:Body></soap12:Envelope>"
)
KR = re.compile(r"<DT>(\d{4}-\d{2}-\d{2})T[^<]*</DT><Rate>([\d.]+)</Rate>")
KEY_RATE_START = date(2013, 9, 13)  # the key rate was introduced on 13 September 2013


def read_inflation(path: Path) -> dict[str, float]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    if not rows or rows[0][0] != "Дата" or "Инфляция" not in str(rows[0][2]):
        raise SourceStructureError(f"{path.name}: unexpected header {rows[0] if rows else None}")
    result = {}
    for row in rows[1:]:
        match = re.match(r"^(\d{2})\.(\d{4})$", str(row[0] or ""))
        if match and isinstance(row[2], int | float):
            result[f"{match.group(2)}-{match.group(1)}"] = float(row[2])
    return result


def read_key_rate(path: Path) -> dict[date, float]:
    rates = {date.fromisoformat(d): float(r) for d, r in KR.findall(path.read_text(encoding="utf-8"))}
    if not rates:
        raise SourceStructureError(f"{path.name}: no <KR> records")
    return rates


def annual_average_calendar_days(rates: dict[date, float], last_day: date) -> dict[str, float]:
    """Average over calendar days; on days without a record the last set rate applies."""
    by_year: dict[int, list[float]] = defaultdict(list)
    current = None
    day = min(rates)
    while day <= last_day:
        current = rates.get(day, current)
        by_year[day.year].append(current)
        day += timedelta(days=1)
    first_full = min(rates).year + (0 if min(rates) == date(min(rates).year, 1, 1) else 1)
    return {str(y): sum(v) / len(v) for y, v in by_year.items() if first_full <= y < last_day.year}


class CbrStats(Connector):
    source_id = "cbr_stats"
    name = "Инфляция и ключевая ставка"
    organization = "Банк России"
    url = "https://www.cbr.ru/hd_base/infl/"
    license = "Материалы сайта Банка России, со ссылкой на источник"
    access = "XLSX-выгрузка и SOAP-веб-сервис DailyInfo"
    role = "actual"
    dedupe = True

    def list_releases(self) -> list[Release]:
        today = date.today()
        return [
            Release(
                release_id=f"cbr_stats_{today:%Y_%m_%d}",
                source_id=self.source_id,
                title=f"Банк России {today:%d.%m.%Y}",
                vintage_date=today,
                url=INFL.format(to=today.strftime("%m%%2F%d%%2F%Y")),
                filename="infl.xlsx",
            )
        ]

    def fetch(self, release: Release) -> Path:
        download(release.url, release.raw_path)
        body = SOAP_BODY.format(start=KEY_RATE_START.isoformat(), end=release.vintage_date.isoformat())
        with client() as http:
            response = http.post(
                SOAP,
                content=body.encode("utf-8"),
                headers={"Content-Type": "application/soap+xml; charset=utf-8"},
            )
            response.raise_for_status()
        (release.raw_path.parent / "keyrate.xml").write_text(response.text, encoding="utf-8")
        return release.raw_path

    def parse(self, raw: Path, release: Release) -> Parsed:
        inflation = read_inflation(raw)
        observations = [
            Observation("cpi_dec", "RU", period[:4], "A", value, "actual")
            for period, value in inflation.items()
            if period.endswith("-12")
        ]
        rates = read_key_rate(raw.parent / "keyrate.xml")
        observations += [
            Observation("key_rate_avg", "RU", year, "A", value, "actual")
            for year, value in annual_average_calendar_days(rates, release.vintage_date).items()
        ]
        return Parsed(observations, data_cutoff=max(inflation) if inflation else None)


register(CbrStats())
