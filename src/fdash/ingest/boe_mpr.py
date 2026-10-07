"""Bank of England MPR: Brent futures assumption (average level in Q4 of each year)."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime
from email.utils import parsedate_to_datetime
from pathlib import Path

from openpyxl import load_workbook

from fdash.ingest.base import Connector, Observation, Parsed, Release, SourceStructureError, client
from fdash.ingest.registry import register

BASE = "https://www.bankofengland.co.uk"
MONTHS = [
    "january",
    "february",
    "march",
    "april",
    "may",
    "june",
    "july",
    "august",
    "september",
    "october",
    "november",
    "december",
]
FIRST_WITH_DATABANK = (2025, 11)
PUBLISHED = re.compile(r"Published on\s*(?:<[^>]+>\s*)*(\d{1,2} [A-Z][a-z]+ \d{4})")
ZIP_LINK = re.compile(r'href="((?:https://www\.bankofengland\.co\.uk)?/-/media/[^"]+\.zip)"')
CUTOFF = re.compile(r"15 working days to (\d{1,2} [A-Z][a-z]+(?: \d{4})?)")


def parse_release_page(html: str) -> tuple[date | None, str] | None:
    """Publication date (None if the page has no 'Published on') and data zip URL;
    None when the report is not out yet (no data archive)."""
    if "to be published" in html.split("</title>", 1)[0]:
        return None
    link = ZIP_LINK.search(html)
    if not link:
        return None
    url = link.group(1)
    published = PUBLISHED.search(html)
    day = datetime.strptime(published.group(1), "%d %B %Y").date() if published else None
    return day, url if url.startswith("http") else BASE + url


QUARTER_NOTE = re.compile(r"average level in (Q[1-4]) of each year")
BLOCKS = ("Latest projections", "Scenario A")
HEADER = re.compile(r"^(\d{4})(?: (Q[1-4]))?$")


def parse_assumptions(rows: list[tuple]) -> tuple[list[tuple[str, float]], str | None]:
    """Oil price assumption as (YYYY-Qn, value) from the central or Scenario A block."""
    note = " ".join(str(r[0]) for r in rows[:3] if r and r[0])
    quarter = QUARTER_NOTE.search(note)
    if not quarter:
        raise SourceStructureError("Other assumptions: 'average level in Qn of each year' note not found")
    footnote = next((str(r[0]) for r in rows if r and str(r[0] or "").startswith("(c)")), "")
    if "Brent futures" not in footnote:
        raise SourceStructureError(f"Other assumptions: unexpected oil footnote: {footnote!r}")
    block_row = start = None
    for i, row in enumerate(rows):
        for c, cell in enumerate(row):
            if isinstance(cell, str) and any(b in cell for b in BLOCKS):
                block_row, start = i, c
                break
        if block_row is not None:
            break
    if block_row is None:
        raise SourceStructureError("Other assumptions: 'Latest projections' / 'Scenario A' block not found")
    header = rows[block_row]
    stop = next((c for c in range(start + 1, len(header)) if header[c]), len(header))
    labels = rows[block_row + 1]
    oil = next((r for r in rows if r and str(r[0] or "").startswith("Oil prices")), None)
    if oil is None:
        raise SourceStructureError("Other assumptions: row 'Oil prices' not found")
    points = []
    for c in range(start, stop):
        match = HEADER.match(str(labels[c]).strip()) if labels[c] is not None else None
        if match and isinstance(oil[c], int | float):
            points.append((f"{match.group(1)}-{match.group(2) or quarter.group(1)}", float(oil[c])))
    cutoff = CUTOFF.search(note)
    return points, cutoff.group(1) if cutoff else None


class BoeMpr(Connector):
    source_id = "boe_mpr"
    name = "Monetary Policy Report: oil price assumption"
    organization = "Банк Англии"
    url = f"{BASE}/monetary-policy-report/monetary-policy-report"
    license = "Bank of England, reuse with attribution"
    access = "ZIP с XLSX по ссылке со страницы выпуска"

    def list_releases(self) -> list[Release]:
        releases = []
        today = date.today()
        with client() as http:
            for year in range(FIRST_WITH_DATABANK[0], today.year + 1):
                for month_index, month in enumerate(MONTHS, start=1):
                    if (year, month_index) < FIRST_WITH_DATABANK or date(year, month_index, 1) > today:
                        continue
                    response = http.get(f"{BASE}/monetary-policy-report/{year}/{month}-{year}")
                    if response.status_code != 200:
                        continue
                    found = parse_release_page(response.text)
                    if not found or (found[0] and found[0] > today):
                        continue
                    published, url = found
                    if published is None:
                        modified = http.head(url).headers.get("last-modified")
                        if not modified:
                            continue
                        published = parsedate_to_datetime(modified).date()
                    releases.append(
                        Release(
                            release_id=f"boe_mpr_{year}_{month_index:02d}",
                            source_id=self.source_id,
                            title=f"MPR {month_index:02d}.{year}",
                            vintage_date=published,
                            url=url,
                            filename=url.rsplit("/", 1)[1],
                        )
                    )
        return releases

    def parse(self, raw: Path, release: Release) -> Parsed:
        with zipfile.ZipFile(raw) as archive:
            name = next((n for n in archive.namelist() if "Projections Databank" in n), None)
            if name is None:
                raise SourceStructureError(f"{raw.name}: Projections Databank not found in the archive")
            wb = load_workbook(io.BytesIO(archive.read(name)), read_only=True, data_only=True)
        sheet = next((n for n in wb.sheetnames if "other assumptions" in n.lower()), None) or next(
            (n for n in wb.sheetnames if n.lower().endswith(". conditioning assumptions")), None
        )
        if sheet is None:
            raise SourceStructureError(f"{name}: assumptions sheet not found")
        rows = list(wb[sheet].iter_rows(min_row=1, max_row=30, max_col=12, values_only=True))
        points, cutoff = parse_assumptions(rows)
        vintage_year = release.vintage_date.year
        observations = [
            Observation(
                "brent_price",
                "WORLD",
                period,
                "Q",
                value,
                "actual" if int(period[:4]) < vintage_year else "forecast",
            )
            for period, value in points
        ]
        return Parsed(observations, data_cutoff=cutoff)


register(BoeMpr())
