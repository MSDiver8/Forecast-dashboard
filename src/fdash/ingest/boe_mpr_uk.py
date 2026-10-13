"""Bank of England MPR: UK CPI inflation and GDP growth projections, history of every report since 2004.

The Projections Databank of the newest report holds one row per past report on sheets
"1. CPI inflation" and "2. GDP growth"; projected cells are shaded (explicit fill or a
conditional format band). The zip is stored once and every release points to it.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange

from fdash import settings
from fdash.ingest.base import (
    Connector,
    Observation,
    Parsed,
    Release,
    SourceStructureError,
    download,
    projection_kind,
)
from fdash.ingest.boe_mpr import BoeMpr
from fdash.ingest.registry import register

SHEETS = {
    "1. CPI inflation": ("cpi_yoy_q", "Four-quarter percentage change in CPI index", "D7G7"),
    "2. GDP growth": ("gdp_yoy_q", "Percentage change on quarter a year earlier", "ABMI"),
}
BAND = re.compile(r"^AND\(COLUMN\(\) >= ROW\(\) - (\d+), COLUMN\(\) <= ROW\(\) \+ (\d+)\)$")
TARGET = re.compile(r"^(\d{4}) Q([1-4])$")
SHADE_THEME = 3  # explicit fill of projected cells
FIRST_MPR = date(2019, 11, 1)  # earlier reports were called Inflation Report
MATCH_DAYS = 45  # a row's month label vs. the exact publication date of a downloaded report


def _bands(sheet) -> list[tuple[CellRange, int, int]]:
    bands = []
    for formatting in sheet.conditional_formatting:
        for rule in formatting.rules:
            formula = (rule.formula or [""])[0]
            match = BAND.match(formula)
            if not match:
                raise SourceStructureError(f"{sheet.title}: unknown projection rule {formula!r}")
            for area in str(formatting.sqref).split():
                bands.append((CellRange(area), int(match.group(1)), int(match.group(2))))
    return bands


def read_history(workbook) -> dict[date, dict[str, list[tuple[str, float, bool]]]]:
    """Report month -> indicator -> (YYYY-Qn, value, projected) for every row of both sheets."""
    history: dict[date, dict[str, list[tuple[str, float, bool]]]] = {}
    for name, (indicator_id, marker, ons_code) in SHEETS.items():
        if name not in workbook.sheetnames:
            raise SourceStructureError(f"sheet {name!r} not found")
        sheet = workbook[name]
        note = " ".join(str(sheet.cell(r, 1).value or "") for r in range(1, 4))
        if marker not in note or ons_code not in note:
            raise SourceStructureError(f"{name}: expected {marker!r} based on {ons_code}, got {note!r}")
        if sheet.cell(5, 1).value != "Date of publication":
            raise SourceStructureError(f"{name}: 'Date of publication' header not in A5")
        targets = {}
        for col in range(2, sheet.max_column + 1):
            match = TARGET.match(str(sheet.cell(5, col).value or ""))
            if match:
                targets[col] = f"{match.group(1)}-Q{match.group(2)}"
        bands = _bands(sheet)
        for row in range(6, sheet.max_row + 1):
            issued = sheet.cell(row, 1).value
            if not isinstance(issued, datetime):
                continue
            points = []
            for col, period in targets.items():
                cell = sheet.cell(row, col)
                if not isinstance(cell.value, int | float) or isinstance(cell.value, bool):
                    continue
                fill = cell.fill
                shaded = (
                    fill.patternType == "solid"
                    and fill.fgColor.type == "theme"
                    and fill.fgColor.theme == SHADE_THEME
                )
                shaded = shaded or any(
                    area.min_row <= row <= area.max_row
                    and area.min_col <= col <= area.max_col
                    and row - before <= col <= row + after
                    for area, before, after in bands
                )
                points.append((period, float(cell.value), shaded))
            if any(projected for _, _, projected in points):
                history.setdefault(issued.date(), {})[indicator_id] = points
    if not history:
        raise SourceStructureError("no report rows with projections")
    return history


@lru_cache(maxsize=2)
def _history_from_zip(path: str, mtime: float) -> dict[date, dict[str, list[tuple[str, float, bool]]]]:
    with zipfile.ZipFile(path) as archive:
        name = next(
            (n for n in archive.namelist() if "Projections Databank" in n and n.endswith(".xlsx")), None
        )
        if name is None:
            raise SourceStructureError(f"{Path(path).name}: Projections Databank not found in the archive")
        workbook = load_workbook(io.BytesIO(archive.read(name)), data_only=True)
    return read_history(workbook)


def history(path: Path) -> dict[date, dict[str, list[tuple[str, float, bool]]]]:
    return _history_from_zip(str(path), path.stat().st_mtime)


def match_exact(label: date, exact: list[date]) -> date | None:
    """Exact publication date of a downloaded report for a row labelled with the report month."""
    near = [d for d in exact if abs((d - label).days) <= MATCH_DAYS]
    return min(near, key=lambda d: abs((d - label).days)) if near else None


class BoeMprUk(Connector):
    source_id = "boe_mpr_uk"
    name = "Monetary Policy Report: UK CPI inflation and GDP growth projections"
    organization = "Банк Англии"
    url = "https://www.bankofengland.co.uk/monetary-policy-report/monetary-policy-report"
    license = "Bank of England, reuse with attribution"
    access = "ZIP с XLSX (Projections Databank) со страницы выпуска; история всех докладов в одном файле"

    def list_releases(self) -> list[Release]:
        reports = BoeMpr().list_releases()  # exact dates and archive links since November 2025
        if not reports:
            return []
        newest = reports[-1]
        shared = f"{newest.vintage_date.isoformat()}/{newest.filename}"
        path = settings.RAW_DIR / self.source_id / shared
        if not path.exists():
            download(newest.url, path)
        exact = [r.vintage_date for r in reports]
        releases = []
        for label in sorted(history(path)):
            published = match_exact(label, exact)
            vintage = published or label
            kind = "MPR" if vintage >= FIRST_MPR else "Inflation Report"
            releases.append(
                Release(
                    release_id=f"{self.source_id}_{vintage:%Y_%m}",
                    source_id=self.source_id,
                    title=f"{kind} {vintage:%m.%Y}" + ("" if published else " (дата ≈)"),
                    vintage_date=vintage,
                    url=newest.url,
                    filename=newest.filename,
                    extra={"row": label.isoformat(), "approximate_date": published is None},
                    shared=shared,
                )
            )
        return releases

    def fetch(self, release: Release) -> Path:
        if not release.raw_path.exists():
            download(release.url, release.raw_path)
        return release.raw_path

    def parse(self, raw: Path, release: Release) -> Parsed:
        rows = history(raw).get(date.fromisoformat(release.extra["row"]))
        if not rows:
            raise SourceStructureError(f"{raw.name}: no row for report {release.extra['row']}")
        observations = [
            Observation(
                indicator_id, "GB", period, "Q", value, projection_kind(period, release.vintage_date, shaded)
            )
            for indicator_id, points in rows.items()
            for period, value, shaded in points
        ]
        actual = [o.target_period for o in observations if o.kind == "actual"]
        return Parsed(observations, data_cutoff=max(actual) if actual else None)


register(BoeMprUk())
