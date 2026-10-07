"""World Bank Commodity Markets Outlook: annual Brent forecast, one XLSX per release."""

from __future__ import annotations

import re
from datetime import date
from email.utils import parsedate_to_datetime
from pathlib import Path

from openpyxl import load_workbook

from fdash.ingest.base import Connector, Observation, Parsed, Release, SourceStructureError, client, http_get
from fdash.ingest.registry import register

PAGE = "https://www.worldbank.org/en/research/commodity-markets/price-forecasts"
FILE = re.compile(r'https://thedocs\.worldbank\.org/[^"]+/CMO-(April|October)-(\d{4})-[Ff]orecasts\.xlsx')
FIRST_WITH_BRENT = (2022, 4)
YEAR_LABEL = re.compile(r"^(\d{4})(f?)$")


def parse_archive(html: str) -> list[tuple[int, int, str]]:
    """(year, month, url) of forecast files that carry the Brent row."""
    result = {}
    for match in FILE.finditer(html):
        month = 4 if match.group(1) == "April" else 10
        year = int(match.group(2))
        if (year, month) >= FIRST_WITH_BRENT:
            result[(year, month)] = match.group(0)
    return sorted((y, m, url) for (y, m), url in result.items())


def parse_forecast_sheet(rows: list[tuple]) -> tuple[list[tuple[str, float, bool]], str | None]:
    header_idx, columns = None, {}
    for idx, row in enumerate(rows[:15]):
        found = {}
        for col, cell in enumerate(row):
            match = YEAR_LABEL.match(str(cell).strip()) if cell is not None else None
            if match and match.group(1) not in found:
                found[match.group(1)] = (col, match.group(2) == "f")
        if len(found) >= 3:
            header_idx, columns = idx, found
            break
    if header_idx is None:
        raise SourceStructureError("Forecast: header row with years not found")
    brent = next(
        (
            row
            for row in rows
            if any(isinstance(c, str) and c.strip().lower() == "crude oil, brent" for c in row[:3])
        ),
        None,
    )
    if brent is None:
        raise SourceStructureError("Forecast: row 'Crude oil, Brent' not found")
    if not any(isinstance(c, str) and c.strip() == "$/bbl" for c in brent):
        raise SourceStructureError("Forecast: Brent row is not in $/bbl")
    points = [
        (year, float(brent[col]), is_forecast)
        for year, (col, is_forecast) in sorted(columns.items())
        if col < len(brent) and isinstance(brent[col], int | float)
    ]
    actual_years = [year for year, _, is_forecast in points if not is_forecast]
    return points, max(actual_years) if actual_years else None


class WbCmo(Connector):
    source_id = "wb_cmo"
    name = "Commodity Markets Outlook"
    organization = "Всемирный банк"
    url = PAGE
    license = "CC BY 4.0"
    access = "XLSX, ссылки со страницы-архива"

    def list_releases(self) -> list[Release]:
        response = http_get(PAGE)
        response.raise_for_status()
        releases = []
        with client() as http:
            for year, month, url in parse_archive(response.text):
                head = http.head(url)
                head.raise_for_status()
                modified = head.headers.get("last-modified")
                published = parsedate_to_datetime(modified).date() if modified else date(year, month, 28)
                releases.append(
                    Release(
                        release_id=f"wb_cmo_{year}_{month:02d}",
                        source_id=self.source_id,
                        title=f"CMO {'апрель' if month == 4 else 'октябрь'} {year}",
                        vintage_date=published,
                        url=url,
                        filename=url.rsplit("/", 1)[1],
                    )
                )
        return releases

    def parse(self, raw: Path, release: Release) -> Parsed:
        wb = load_workbook(raw, read_only=True, data_only=True)
        if "Forecast" not in wb.sheetnames:
            raise SourceStructureError(f"{raw.name}: sheet Forecast not found")
        rows = list(wb["Forecast"].iter_rows(max_col=40, values_only=True))
        points, cutoff = parse_forecast_sheet(rows)
        observations = [
            Observation("brent_price", "WORLD", year, "A", value, "forecast" if is_forecast else "estimate")
            for year, value, is_forecast in points
        ]
        return Parsed(observations, data_cutoff=cutoff)


register(WbCmo())
