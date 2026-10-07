"""Bank of Russia macroeconomic survey: median and 10-90% range per survey, full history in one XLSX."""

from __future__ import annotations

import shutil
import tempfile
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from fdash.ingest.base import Connector, Observation, Parsed, Release, SourceStructureError, download
from fdash.ingest.registry import register

PAGE = "https://www.cbr.ru/statistics/ddkp/mo_br/"
FILE = "https://www.cbr.ru/Content/Document/File/144490/full.xlsx"
SHEETS = {
    "1": ("cpi_dec", "ИПЦ"),
    "2": ("cpi_avg", "в среднем за год"),
    "3": ("key_rate_avg", "Ключевая ставка"),
    "4": ("gdp_growth", "ВВП"),
    "5": ("unemployment_avg", "в среднем за год"),
    "9": ("exports_usd", "Экспорт"),
    "10": ("imports_usd", "Импорт"),
    "12": ("oil_tax_price", "налогообложения"),
    "13": ("brent_price", "Brent"),
}
AREAS = {"brent_price": "WORLD"}
BLOCKS = {"median": "Медиана", "p90": "90-й процентиль", "p10": "10-й процентиль"}
MONTHS_RU = [
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
]


def _num(value) -> float | None:
    return float(value) if isinstance(value, int | float) and not isinstance(value, bool) else None


def read_sheet(rows: list[tuple]) -> dict[date, dict[str, dict[str, float | None]]]:
    """survey month -> target year -> {median, p10, p90}."""
    header = rows[5]
    survey_cols = {c: v.date() for c, v in enumerate(header) if isinstance(v, datetime) and c >= 4}
    if not survey_cols:
        raise SourceStructureError("survey dates not found in row 6")
    starts = {}
    for i, row in enumerate(rows):
        label = str(row[1] or "").strip()
        for key, name in BLOCKS.items():
            if label == name:
                starts[key] = i
    missing = set(BLOCKS) - set(starts)
    if missing:
        raise SourceStructureError(f"blocks not found: {sorted(missing)}")
    result: dict[date, dict[str, dict[str, float | None]]] = {d: {} for d in survey_cols.values()}
    for key, start in starts.items():
        for row in rows[start : start + 12]:
            target = row[3]
            if not isinstance(target, datetime):
                if row is not rows[start] and str(row[1] or "").strip():
                    break
                continue
            for col, survey in survey_cols.items():
                value = _num(row[col]) if col < len(row) else None
                result[survey].setdefault(str(target.year), {})[key] = value
    return result


@lru_cache(maxsize=4)
def _workbook_tables(path: str, mtime: float) -> dict[str, dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    tables = {}
    for sheet, (indicator_id, marker) in SHEETS.items():
        if sheet not in wb.sheetnames:
            raise SourceStructureError(f"sheet {sheet} not found")
        rows = list(wb[sheet].iter_rows(min_row=1, max_row=95, values_only=True))
        title = str(rows[3][1] or "")
        if marker not in title:
            raise SourceStructureError(f"sheet {sheet}: expected '{marker}' in title, got {title!r}")
        tables[indicator_id] = read_sheet(rows)
    return tables


class CbrSurvey(Connector):
    source_id = "cbr_survey"
    name = "Макроэкономический опрос"
    organization = "Банк России"
    url = PAGE
    license = "Материалы сайта Банка России, со ссылкой на источник"
    access = "XLSX по стабильной ссылке, вся история в одном файле"

    def __init__(self) -> None:
        self._latest: Path | None = None

    def _download_latest(self) -> Path:
        if self._latest is None or not self._latest.exists():
            tmp = Path(tempfile.mkdtemp(prefix="fdash_cbr_")) / "full.xlsx"
            download(FILE, tmp)
            self._latest = tmp
        return self._latest

    def list_releases(self) -> list[Release]:
        latest = self._download_latest()
        tables = _workbook_tables(str(latest), latest.stat().st_mtime)
        surveys = sorted(next(iter(tables.values())).keys())
        since = date.today() - timedelta(days=365 * 5 + 31)
        return [
            Release(
                release_id=f"cbr_survey_{s:%Y_%m}",
                source_id=self.source_id,
                title=f"Опрос {MONTHS_RU[s.month - 1]} {s.year} (дата ≈)",
                vintage_date=s,
                url=FILE,
                filename="full.xlsx",
                extra={"survey_month": s.isoformat(), "approximate_date": True},
            )
            for s in surveys
            if s >= since
        ]

    def fetch(self, release: Release) -> Path:
        release.raw_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(self._download_latest(), release.raw_path)
        return release.raw_path

    def parse(self, raw: Path, release: Release) -> Parsed:
        survey = date.fromisoformat(release.extra["survey_month"])
        tables = _workbook_tables(str(raw), raw.stat().st_mtime)
        observations = []
        for indicator_id, table in tables.items():
            for year, stats in table.get(survey, {}).items():
                if stats.get("median") is None:
                    continue
                observations.append(
                    Observation(
                        indicator_id,
                        AREAS.get(indicator_id, "RU"),
                        year,
                        "A",
                        stats["median"],
                        "estimate" if int(year) < survey.year else "forecast",
                        lower=stats.get("p10"),
                        upper=stats.get("p90"),
                    )
                )
        return Parsed(observations, data_cutoff=None)


register(CbrSurvey())
