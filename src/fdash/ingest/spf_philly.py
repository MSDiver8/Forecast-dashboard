"""Philadelphia Fed Survey of Professional Forecasters: median forecast of US real GDP growth, quarterly.

mediangrowth.xlsx holds every survey since 1968 (sheet RGDP: drgdp2 is the quarter of the survey,
drgdp3..6 the four following quarters; annualized quarter-over-quarter growth of the median level).
News release dates come from spf-release-dates.txt.
"""

from __future__ import annotations

import re
import shutil
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from fdash.ingest.base import (
    Connector,
    Observation,
    Parsed,
    Release,
    SourceStructureError,
    copy_into,
    download_temp,
)
from fdash.ingest.registry import register

# "In the survey of 1992:Q1, the survey's measure of output switches from GNP to GDP" (SPF documentation).
FIRST_GDP_SURVEY = (1992, 1)

ASSETS = (
    "https://www.philadelphiafed.org/-/media/frbp/assets/surveys-and-data/survey-of-professional-forecasters/"
)
FILE = ASSETS + "historical-data/mediangrowth.xlsx"
DATES = ASSETS + "spf-release-dates.txt"
HEADER = ("YEAR", "QUARTER", "drgdp2", "drgdp3", "drgdp4", "drgdp5", "drgdp6")
DATE_LINE = re.compile(r"^\s*(\d{4})?\s+Q([1-4])\s+(\d{1,2}/\d{1,2}/\d{2})\**\s+(\d{1,2}/\d{1,2}/\d{2})\**")


def release_dates(text: str) -> dict[tuple[int, int], date]:
    """(year, quarter) -> news release date. The year is written only on the Q1 line of each year."""
    result, year = {}, None
    for line in text.splitlines():
        match = DATE_LINE.match(line)
        if not match:
            continue
        if match.group(1):
            year = int(match.group(1))
        if year is None:
            continue
        result[(year, int(match.group(2)))] = datetime.strptime(match.group(4), "%m/%d/%y").date()
    if not result:
        raise SourceStructureError("spf-release-dates.txt: no survey lines found")
    return result


def read_rgdp(rows: list[tuple]) -> dict[tuple[int, int], list[float | None]]:
    """(year, quarter) of the survey -> growth forecasts for the survey quarter and four next ones."""
    if tuple(rows[0][: len(HEADER)]) != HEADER:
        raise SourceStructureError(f"RGDP: header {HEADER} expected, got {rows[0]}")
    result = {}
    for row in rows[1:]:
        if not isinstance(row[0], int | float) or not isinstance(row[1], int | float):
            continue
        values = [float(v) if isinstance(v, int | float) else None for v in row[2:7]]
        result[(int(row[0]), int(row[1]))] = values
    return result


@lru_cache(maxsize=2)
def _surveys(path: str, mtime: float) -> dict[tuple[int, int], list[float | None]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if "RGDP" not in wb.sheetnames:
        raise SourceStructureError(f"{Path(path).name}: sheet RGDP not found")
    return read_rgdp(list(wb["RGDP"].iter_rows(values_only=True)))


def surveys(path: Path) -> dict[tuple[int, int], list[float | None]]:
    return _surveys(str(path), path.stat().st_mtime)


class SpfPhilly(Connector):
    source_id = "spf_philly"
    name = "Survey of Professional Forecasters: real GDP (median)"
    organization = "ФРБ Филадельфии"
    url = "https://www.philadelphiafed.org/surveys-and-data/real-time-data-research/survey-of-professional-forecasters"
    license = "Federal Reserve Bank of Philadelphia, free use with attribution"
    access = "XLSX по стабильной ссылке, вся история опросов в одном файле; даты выхода — TXT"

    def __init__(self) -> None:
        self._latest: tuple[Path, Path] | None = None

    def _download_latest(self) -> tuple[Path, Path]:
        if self._latest is None or not self._latest[0].exists():
            self._latest = (
                download_temp(FILE, "mediangrowth.xlsx"),
                download_temp(DATES, "spf-release-dates.txt"),
            )
        return self._latest

    def list_releases(self) -> list[Release]:
        workbook, dates_file = self._download_latest()
        dates = release_dates(dates_file.read_text(encoding="utf-8", errors="replace"))
        releases = []
        for year, quarter in sorted(surveys(workbook)):
            if (year, quarter) < FIRST_GDP_SURVEY:
                continue  # real GNP before 1992 Q1: another definition than the BEA real GDP actuals
            published = dates.get((year, quarter))
            if published is None:
                continue
            releases.append(
                Release(
                    release_id=f"{self.source_id}_{year}_q{quarter}",
                    source_id=self.source_id,
                    title=f"SPF {quarter} кв. {year}",
                    vintage_date=published,
                    url=FILE,
                    filename="mediangrowth.xlsx",
                    extra={"survey": f"{year}-Q{quarter}"},
                )
            )
        return releases

    def fetch(self, release: Release) -> Path:
        workbook, dates_file = self._download_latest()
        path = copy_into(workbook, release)
        shutil.copyfile(dates_file, release.release_dir / dates_file.name)
        return path

    def parse(self, raw: Path, release: Release) -> Parsed:
        survey = release.extra["survey"]
        year, quarter = int(survey[:4]), int(survey[-1])
        values = surveys(raw).get((year, quarter))
        if values is None:
            raise SourceStructureError(f"{raw.name}: survey {survey} not found")
        observations = []
        for step, value in enumerate(values):
            if value is None:
                continue
            index = year * 4 + quarter - 1 + step
            period = f"{index // 4:04d}-Q{index % 4 + 1}"
            observations.append(Observation("gdp_qoq_saar", "US", period, "Q", value, "forecast"))
        return Parsed(observations, data_cutoff=None)


register(SpfPhilly())
